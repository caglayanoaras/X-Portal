import os
import openpyxl

def process_excel(input_path, key_value_path, target_column_name):
    """
    Reads a key-value mapping from key_value.xlsx, searches for a specific 
    column in the first 10 rows of input_path, maps the values, and saves 
    the result to a new file.
    """
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Input file not found: {input_path}")
    if not os.path.exists(key_value_path):
        raise FileNotFoundError(f"Key-value file not found: {key_value_path}")

    print(f"Loading key-value mapping from '{key_value_path}'...")
    
    wb_kv = openpyxl.load_workbook(key_value_path, data_only=True)
    ws_kv = wb_kv.active
    
    mapping = {}
    # Read the mapping. We assume column A (0) is ID, column B (1) is Name
    for row in ws_kv.iter_rows(min_row=1, values_only=True):
        if len(row) >= 2:
            key, val = row[0], row[1]
            if key is not None:
                mapping[key] = val
                # Also store a string version of the key to handle 
                # type mismatches (e.g., int 123 vs string "123")
                mapping[str(key)] = val

    wb_kv.close()
    print(f"Loaded {len(mapping)} mapping rules (including string fallbacks).")

    print(f"Loading input file '{input_path}'...")
    wb_input = openpyxl.load_workbook(input_path)
    ws_input = wb_input.active

    print(f"Searching for column '{target_column_name}' in the first 10 rows...")
    matches = []
    
    # Search space: rows 1 to 10, all columns
    max_search_row = min(ws_input.max_row, 10)
    for r in range(1, max_search_row + 1):
        for c in range(1, ws_input.max_column + 1):
            cell_val = ws_input.cell(row=r, column=c).value
            # Check if the cell value exactly matches the target column name
            if isinstance(cell_val, str) and cell_val.strip() == target_column_name:
                matches.append((r, c))

    if len(matches) == 0:
        raise ValueError(f"Error: Column name '{target_column_name}' was not found in the first 10 rows.")
    elif len(matches) > 1:
        match_locations = [f"Row {m[0]}, Col {m[1]}" for m in matches]
        raise ValueError(f"Error: Column name '{target_column_name}' found multiple times: {', '.join(match_locations)}. Please ensure it only appears once.")

    target_row, target_col = matches[0]
    print(f"Found column '{target_column_name}' at Row {target_row}, Column {target_col}.")

    # Insert a new column at the target column index.
    # This pushes the original column to the right (target_col + 1).
    # The new empty column takes the place of target_col.
    ws_input.insert_cols(target_col)

    # Set the headers
    # New column gets the original target_column_name
    ws_input.cell(row=target_row, column=target_col).value = target_column_name
    
    # Original column gets the "__" suffix
    old_col_header = f"{target_column_name}__"
    ws_input.cell(row=target_row, column=target_col + 1).value = old_col_header

    print("Applying mapping to rows...")
    # Start iterating just beneath the header row
    for r in range(target_row + 1, ws_input.max_row + 1):
        # Read the old value from the shifted original column
        old_val = ws_input.cell(row=r, column=target_col + 1).value
        
        if old_val is None:
            # Leave empty cells as is
            ws_input.cell(row=r, column=target_col).value = None
        else:
            # Look up the mapped value. First try the exact type, then try string cast.
            # If neither exists, fallback to the original old_val (leave as is).
            mapped_val = mapping.get(old_val, mapping.get(str(old_val), old_val))
            
            # Write the mapped value to the new column
            ws_input.cell(row=r, column=target_col).value = mapped_val

    # Construct output filename: append "__" before the extension
    base_name, ext = os.path.splitext(input_path)
    output_path = f"{base_name}__{ext}"
    
    print(f"Saving changes to '{output_path}'...")
    wb_input.save(output_path)
    wb_input.close()
    print("Done!")

if __name__ == "__main__":
    # ==========================================
    # USER CONFIGURATION
    # ==========================================
    INPUT_EXCEL_PATH = "input_data.xlsx"         # Change this to your input file path
    KEY_VALUE_EXCEL_PATH = "key_value.xlsx"      # Change this to your key_value file path
    COLUMN_TO_FIND = "MyColumn"                  # Change this to the column name you want to find
    # ==========================================
    
    try:
        # Create dummy files for demonstration if they don't exist
        # (You can remove this block when running with real files)
        if not os.path.exists(KEY_VALUE_EXCEL_PATH):
            print("Creating a sample key_value.xlsx for testing...")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append([1, "Apple"])
            ws.append([2, "Banana"])
            ws.append(["XYZ", "Orange"])
            wb.save(KEY_VALUE_EXCEL_PATH)
            
        if not os.path.exists(INPUT_EXCEL_PATH):
            print("Creating a sample input_data.xlsx for testing...")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Redundant Row 1", "", ""])
            ws.append(["Redundant Row 2", "", ""])
            ws.append(["ID", "MyColumn", "Other Data"]) # Header is on Row 3
            ws.append([101, 1, "Data A"])
            ws.append([102, "XYZ", "Data B"])
            ws.append([103, 999, "Data C"]) # 999 has no mapping, will remain 999
            ws.append([104, None, "Data D"]) # Empty, will remain empty
            wb.save(INPUT_EXCEL_PATH)
            
        process_excel(
            input_path=INPUT_EXCEL_PATH, 
            key_value_path=KEY_VALUE_EXCEL_PATH, 
            target_column_name=COLUMN_TO_FIND
        )
        
    except Exception as e:
        print(f"Script stopped due to an error: {e}")